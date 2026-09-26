"""
Lógica de análisis y corrección APA 7.ª edición exclusiva para archivos .docx.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, asdict
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

APA_FONTS = {
    ("Times New Roman", 12.0),
    ("Arial", 11.0),
    ("Calibri", 11.0),
    ("Georgia", 11.0),
    ("Lucida Sans Unicode", 10.0),
}


@dataclass
class Check:
    category: str
    title: str
    status: str
    detail: str
    recommendation: str = ""

    def to_dict(self) -> dict:
        return asdict(self)


def iter_paragraphs(document: Document) -> Iterable:
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def non_empty_paragraphs(document: Document) -> list:
    return [p for p in iter_paragraphs(document) if p.text.strip()]


def length_cm(value) -> float | None:
    if value is None:
        return None
    try:
        return round(float(value.cm), 2)
    except (AttributeError, TypeError, ValueError):
        return None


def paragraph_font_samples(paragraphs: list) -> list[tuple[str, float]]:
    samples: list[tuple[str, float]] = []
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            font_name = run.font.name or "Times New Roman"
            size = run.font.size.pt if run.font.size else 12.0
            samples.append((str(font_name), round(float(size), 1)))
    return samples


def has_page_field(document: Document) -> bool:
    for section in document.sections:
        try:
            header_xml = section.header._element.xml
            if "PAGE" in header_xml.upper():
                return True
        except Exception:
            continue
    return False


def is_reference_heading(text: str) -> bool:
    text_clean = text.strip()
    text_clean = re.sub(r"^\d+[\.\)]?\s*|\s*\d+$", "", text_clean)
    pattern = (
        r"^\s*(referencias(\s+bibliográficas|\s+bibliograficas)?|"
        r"references|bibliografía|bibliografia)\b"
    )
    return bool(re.search(pattern, text_clean, re.I))


def find_reference_start(paragraphs: list) -> tuple[int | None, str]:
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip() if hasattr(paragraph, "text") else str(paragraph).strip()
        if is_reference_heading(text):
            return index, text
    return None, ""


def first_author_surname(author: str) -> str:
    if not isinstance(author, str):
        return ""
    cleaned = re.sub(r"\bet\s+al\.?\b", "", author, flags=re.I)
    cleaned = re.split(r"\s+(?:y|e|and|&)\s+", cleaned, maxsplit=1, flags=re.I)[0]
    cleaned = cleaned.split(",", 1)[0].strip().lower()
    return re.sub(r"[^a-záéíóúüñ0-9 -]", "", cleaned).strip()
    
def extract_citations(text: str) -> list[dict]:
    citations: list[dict] = []

    parenthetical_groups = re.findall(
        r"\(([^()]{3,200}?(?:19|20)\d{2}[a-z]?(?:[^()]{0,80})?)\)",
        text,
    )

    for group in parenthetical_groups:
        for part in re.split(r"\s*;\s*", group):
            match = re.search(
                r"(?P<autor>[A-ZÁÉÍÓÚÑ][^()]{1,120}?),\s*"
                r"(?P<anio>(?:19|20)\d{2}[a-z]?)\b",
                part,
                re.I,
            )
            if not match:
                continue

            autor = match.group("autor").strip()
            anio = match.group("anio").strip()

            if len(autor) < 2:
                continue

            citations.append(
                {
                    "autor": autor,
                    "anio": anio,
                    "tipo": "parentética",
                    "texto": f"({part.strip()})",
                }
            )

    narrative = re.findall(
        r"\b([A-ZÁÉÍÓÚÑ][A-Za-zÁÉÍÓÚÑáéíóúüñ'-]{1,}"
        r"(?:\s+(?:y|e|&|and)\s+[A-ZÁÉÍÓÚÑ][A-Za-zÁÉÍÓÚÑáéíóúüñ'-]{1,})?"
        r"(?:\s+et\s+al\.?)?)"
        r"\s+\(((?:19|20)\d{2}[a-z]?)\)",
        text,
    )

    for autor, anio in narrative:
        citations.append(
            {
                "autor": autor.strip(),
                "anio": anio.strip(),
                "tipo": "narrativa",
                "texto": f"{autor.strip()} ({anio})",
            }
        )

    unique = []
    seen = set()
    for citation in citations:
        key = (first_author_surname(citation["autor"]), str(citation["anio"]))
        if key not in seen and key[0]:
            seen.add(key)
            unique.append(citation)

    return unique


def extract_reference_entries(reference_paragraphs: list) -> list[dict]:
    entries: list[dict] = []
    for paragraph in reference_paragraphs:
        text = paragraph.text.strip() if hasattr(paragraph, "text") else str(paragraph).strip()
        if not text:
            continue

        year_match = re.search(r"\b((?:19|20)\d{2}[a-z]?)\b", text)
        anio_str = year_match.group(1) if year_match else ""

        if "," in text:
            author = text.split(",", 1)[0].strip()
        else:
            author = text.split(".", 1)[0].strip()

        entries.append(
            {
                "autor": author,
                "anio": anio_str,
                "texto": text,
                "surname": first_author_surname(author),
            }
        )
    return entriesdef analyze_document(data: bytes, filename: str = "") -> dict:
    if not filename.lower().endswith(".docx"):
        raise ValueError("La aplicación está optimizada exclusivamente para archivos .docx.")

    document = Document(io.BytesIO(data))
    paragraphs = non_empty_paragraphs(document)

    reference_start, raw_heading_text = find_reference_start(paragraphs)
    body_paragraphs = paragraphs[:reference_start] if reference_start is not None else paragraphs
    reference_paragraphs = (
        paragraphs[reference_start + 1 :] if reference_start is not None else []
    )

    checks: list[Check] = []

    # 1. Márgenes
    section = document.sections[0] if document.sections else None
    if section:
        top_cm = length_cm(getattr(section, "top_margin", None))
        bottom_cm = length_cm(getattr(section, "bottom_margin", None))
        left_cm = length_cm(getattr(section, "left_margin", None))
        right_cm = length_cm(getattr(section, "right_margin", None))

        margin_errors = []
        for name, value in [
            ("Superior", top_cm),
            ("Inferior", bottom_cm),
            ("Izquierdo", left_cm),
            ("Derecho", right_cm),
        ]:
            if value is None or abs(value - 2.54) > 0.1:
                margin_errors.append(
                    f"{name} ({value} cm)" if value is not None else f"{name} (no definido)"
                )

        if not margin_errors:
            checks.append(
                Check(
                    "Formato",
                    "Márgenes de 2,54 cm",
                    "ok",
                    "Los cuatro márgenes están configurados correctamente a 2,54 cm (1 pulgada).",
                )
            )
        else:
            checks.append(
                Check(
                    "Formato",
                    "Márgenes de 2,54 cm",
                    "error",
                    f"Se encontraron márgenes incorrectos en: {', '.join(margin_errors)}.",
                    "Ajusta los márgenes a 2,54 cm en Disposición > Márgenes de Word.",
                )
            )

    # 2. Tipografía
    font_samples = paragraph_font_samples(paragraphs)
    font_ok = not font_samples or all(s in APA_FONTS for s in font_samples)
    font_label = ", ".join(f"{name} {size:g}" for name, size in sorted(set(font_samples))[:4])

    checks.append(
        Check(
            "Formato",
            "Tipografía legible y consistente",
            "ok" if font_ok else "warning",
            f"Se detectó: {font_label or 'tipografía heredada'}."
            if font_ok
            else "Hay tipografías que no coinciden con las opciones estándar de APA 7.",
            "Usa Times New Roman 12 pt, Arial 11 pt, Calibri 11 pt o Georgia 11 pt.",
        )
    )

    # 3. Interlineado
    spacing_values = []
    for paragraph in body_paragraphs:
        value = paragraph.paragraph_format.line_spacing
        if isinstance(value, (int, float)):
            spacing_values.append(float(value))

    spacing_ok = not spacing_values or all(abs(v - 2.0) <= 0.05 for v in spacing_values)
    checks.append(
        Check(
            "Formato",
            "Interlineado doble",
            "ok" if spacing_ok else "warning",
            "Los párrafos del cuerpo usan interlineado doble (2.0)."
            if spacing_ok
            else "Hay párrafos con interlineado distinto de 2,0.",
            "Aplica interlineado doble a todo el texto del documento.",
        )
    )

    # 4. Sangría de primera línea
    body_without_headings = [
        p
        for p in body_paragraphs[1:]
        if not re.match(
            r"^(resumen|abstract|introducción|introduccion|método|metodo|resultados|"
            r"discusión|discusion|conclusión|conclusion)$",
            p.text.strip(),
            re.I,
        )
    ]

    indent_values = [
        length_cm(p.paragraph_format.first_line_indent)
        for p in body_without_headings
        if p.paragraph_format.first_line_indent is not None
    ]

    indent_ok = not indent_values or (
        sum(abs(v - 1.27) <= 0.15 for v in indent_values if v is not None)
        >= max(1, int(len(indent_values) * 0.65))
    )

    checks.append(
        Check(
            "Formato",
            "Sangría de primera línea",
            "ok" if indent_ok else "warning",
            "La mayoría de los párrafos del cuerpo tienen sangría de 1,27 cm."
            if indent_ok
            else "La sangría de primera línea es irregular o está ausente.",
            "Aplica sangría de primera línea de 1,27 cm a cada párrafo del cuerpo.",
        )
    )

    # 5. Numeración de página
    page_status = "ok" if has_page_field(document) else "warning"
    checks.append(
        Check(
            "Formato",
            "Numeración de páginas",
            page_status,
            "El documento contiene numeración de páginas en el encabezado."
            if page_status == "ok"
            else "No se detectó el campo de numeración de páginas en el encabezado.",
            "Añade el número de página alineado a la derecha en el encabezado superior.",
        )
    )# 6. Sección de referencias
    if reference_start is None:
        checks.append(
            Check(
                "Referencias",
                "Sección de referencias",
                "error",
                "No se encontró un encabezado de referencias al final del documento.",
                "Se agregará automáticamente la página de Referencias con el formato correcto al descargar el archivo corregido.",
            )
        )
    else:
        heading_clean = raw_heading_text.strip()
        if heading_clean.lower() in ("referencias", "references"):
            checks.append(
                Check(
                    "Referencias",
                    "Sección de referencias",
                    "ok",
                    "Se encontró la sección «Referencias» con la titulación correcta.",
                )
            )
        else:
            checks.append(
                Check(
                    "Referencias",
                    "Sección de referencias",
                    "warning",
                    f"Se detectó la sección bajo el título «{heading_clean}».",
                    "En APA 7.ª edición, el título oficial debe ser exactamente «Referencias».",
                )
            )

    # 7. Citas y correspondencia
    full_text = "\n".join(p.text for p in body_paragraphs)
    citation_matches = extract_citations(full_text)
    citation_count = len(citation_matches)

    reference_entries = extract_reference_entries(reference_paragraphs)
    citation_keys = {(first_author_surname(c["autor"]), str(c["anio"])) for c in citation_matches}
    reference_keys = {(e["surname"], str(e["anio"])) for e in reference_entries}

    citations_without_reference = [
        c
        for c in citation_matches
        if (first_author_surname(c["autor"]), str(c["anio"])) not in reference_keys
    ]

    uncited_references = [
        e
        for e in reference_entries
        if not e["anio"] or (e["surname"], str(e["anio"])) not in citation_keys
    ]

    checks.append(
        Check(
            "Citas",
            "Citas dentro del texto",
            "ok" if citation_count > 0 else "warning",
            f"Se detectaron {citation_count} cita(s) con formato autor-año."
            if citation_count
            else "No se detectaron patrones claros de cita autor-año en el texto.",
            "Asegúrate de que cada afirmación tomada de una fuente externa tenga su cita correspondiente.",
        )
    )

    if citation_count > 0 and reference_start is not None:
        if not citations_without_reference:
            checks.append(
                Check(
                    "Citas",
                    "Correspondencia citas-referencias",
                    "ok",
                    "Todas las citas encontradas tienen una entrada coincidente en la lista de referencias.",
                )
            )
        else:
            unmatched_labels = ", ".join(
                f"{c['autor']}, {c['anio']}" for c in citations_without_reference[:5]
            )
            extra = len(citations_without_reference) - 5
            checks.append(
                Check(
                    "Citas",
                    "Correspondencia citas-referencias",
                    "warning",
                    f"Hay {len(citations_without_reference)} cita(s) sin una referencia coincidente: {unmatched_labels}"
                    + (f" y {extra} más." if extra > 0 else "."),
                    "Verifica que cada cita en el texto tenga su correspondiente entrada en la lista de referencias.",
                )
            )

    reference_count = len(reference_paragraphs)
    checks.append(
        Check(
            "Referencias",
            "Entradas bibliográficas",
            "ok" if reference_count > 0 else ("error" if reference_start is None else "warning"),
            f"Se detectaron {reference_count} entrada(s) en la sección de referencias."
            if reference_count
            else "La lista de referencias está vacía o no se ha creado.",
            "Incluye la ficha bibliográfica completa para cada fuente citada.",
        )
    )

    return {
        "checks": [c.to_dict() for c in checks],
        "paragraph_count": len(paragraphs),
        "word_count": len(re.findall(r"\b[\wÁÉÍÓÚÜÑáéíóúüñ'-]+\b", full_text)),
        "citation_count": citation_count,
        "citations": citation_matches,
        "reference_count": reference_count,
        "references": reference_entries,
        "citations_without_reference": citations_without_reference,
        "uncited_references": uncited_references,
        "ok_count": sum(c.status == "ok" for c in checks),
        "issue_count": sum(c.status != "ok" for c in checks),
        "is_pdf": False,
    }def set_run_font(run, name: str = "Times New Roman", size: int = 12) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:cs"), name)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "1"

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=12)


def correct_document(data: bytes, filename: str = "") -> bytes:
    if not filename.lower().endswith(".docx"):
        raise ValueError("Solo se pueden corregir archivos .docx.")

    document = Document(io.BytesIO(data))

    # Márgenes y numeración de página
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        header = section.header
        header_paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        if "PAGE" not in header._element.xml.upper():
            header_paragraph.clear()
            add_page_number(header_paragraph)

    # Estilo Normal
    try:
        normal_style = document.styles["Normal"]
        normal_style.font.name = "Times New Roman"
        normal_style.font.size = Pt(12)
        if normal_style._element.rPr is not None and normal_style._element.rPr.rFonts is not None:
            normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        normal_style.paragraph_format.line_spacing = 2
        normal_style.paragraph_format.space_after = Pt(0)
    except KeyError:
        pass

    paragraphs = non_empty_paragraphs(document)
    reference_start, _ = find_reference_start(paragraphs)
    reference_mode = False

    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()

        # Encabezado de Referencias
        if reference_start is not None and index == reference_start:
            reference_mode = True
            paragraph.text = "Referencias"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.left_indent = Inches(0)
            paragraph.paragraph_format.line_spacing = 2
            for run in paragraph.runs:
                set_run_font(run)
                run.bold = True
            continue

        # Formato general de párrafo
        paragraph.paragraph_format.line_spacing = 2
        paragraph.paragraph_format.space_after = Pt(0)

        for run in paragraph.runs:
            if not run.font.name or run.font.name not in {
                "Times New Roman", "Arial", "Calibri", "Georgia", "Lucida Sans Unicode"
            }:
                set_run_font(run)
            elif run.font.size is None:
                run.font.size = Pt(12)

        is_heading = (
            (paragraph.style and paragraph.style.name and paragraph.style.name.lower().startswith("heading"))
            or text.lower()
            in {
                "resumen",
                "abstract",
                "introducción",
                "introduccion",
                "método",
                "metodo",
                "resultados",
                "discusión",
                "discusion",
                "conclusión",
                "conclusion",
            }
        )

        if is_heading:
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.left_indent = Inches(0)
            continue

        if reference_mode:
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        else:
            paragraph.paragraph_format.left_indent = Inches(0)
            paragraph.paragraph_format.first_line_indent = Inches(0.5)# Si no existía sección de referencias, la creamos
    if reference_start is None:
        document.add_page_break()

        ref_heading = document.add_paragraph()
        ref_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = ref_heading.add_run("Referencias")
        set_run_font(run)
        run.bold = True
        ref_heading.paragraph_format.line_spacing = 2
        ref_heading.paragraph_format.first_line_indent = Inches(0)
        ref_heading.paragraph_format.left_indent = Inches(0)

        full_text = "\n".join(p.text for p in paragraphs)
        citations = extract_citations(full_text)

        if citations:
            sorted_authors = sorted(set(c["autor"] for c in citations if c["autor"]))
            for autor in sorted_authors:
                p = document.add_paragraph()
                p.paragraph_format.line_spacing = 2
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)

                run_entry = p.add_run(f"{autor}. (Año). ")
                set_run_font(run_entry)

                run_title = p.add_run("[Título del documento o publicación en cursiva]. ")
                set_run_font(run_title)
                run_title.italic = True

                run_source = p.add_run("[Nombre de la fuente, Editorial o URL].")
                set_run_font(run_source)
        else:
            p = document.add_paragraph()
            p.paragraph_format.line_spacing = 2
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            run_empty = p.add_run(
                "[Añade aquí las fuentes bibliográficas citadas ordenadas alfabéticamente]."
            )
            set_run_font(run_empty)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
